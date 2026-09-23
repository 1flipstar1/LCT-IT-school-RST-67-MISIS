"""Safe first-sheet parsers for legacy XLS and modern XLSX workbooks."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import xlrd

from app.core.config import settings
from app.core.errors import APIError


def _cell_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _check_size(rows: list[list[Any]]) -> list[list[Any]]:
    if len(rows) > settings.max_import_rows:
        raise APIError(422, "import_too_many_rows", f"В файле больше {settings.max_import_rows} строк.")
    width = max((len(row) for row in rows), default=0)
    if width > settings.max_import_columns:
        raise APIError(422, "import_too_many_columns", f"В файле больше {settings.max_import_columns} колонок.")
    while rows and all(str(value).strip() == "" for value in rows[-1]):
        rows.pop()
    if len(rows) < 2:
        raise APIError(422, "import_empty", "В файле нет строк с данными.")
    return rows


def _parse_xlsx(data: bytes) -> list[list[Any]]:
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        if sheet.max_row > settings.max_import_rows or sheet.max_column > settings.max_import_columns:
            workbook.close()
            raise APIError(422, "import_dimensions_exceeded", "Размер таблицы превышает допустимый предел.")
        rows = [[_cell_value(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return _check_size(rows)
    except APIError:
        raise
    except Exception as exc:
        raise APIError(422, "invalid_xlsx", "Не удалось прочитать XLSX-файл.") from exc


def _parse_xls(data: bytes) -> list[list[Any]]:
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows > settings.max_import_rows or sheet.ncols > settings.max_import_columns:
            workbook.release_resources()
            raise APIError(422, "import_dimensions_exceeded", "Размер таблицы превышает допустимый предел.")
        rows: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            values: list[Any] = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(cell.value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    value = ""
                else:
                    value = cell.value
                values.append(_cell_value(value))
            rows.append(values)
        workbook.release_resources()
        return _check_size(rows)
    except APIError:
        raise
    except Exception as exc:
        raise APIError(422, "invalid_xls", "Не удалось прочитать XLS-файл.") from exc


def parse_workbook(data: bytes, filename: str) -> list[list[Any]]:
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        return _parse_xlsx(data)
    if extension == ".xls":
        return _parse_xls(data)
    raise APIError(415, "unsupported_import_format", "Поддерживаются только файлы XLS и XLSX.")
