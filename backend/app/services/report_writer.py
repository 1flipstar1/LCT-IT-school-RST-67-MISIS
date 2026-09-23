"""Binary report writers used by the reports worker."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import xlwt
from fpdf import FPDF, FontFace
from openpyxl import Workbook

from app.core.config import BACKEND_DIR, PROJECT_DIR
from app.core.errors import APIError


CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "pdf": "application/pdf",
}


def _xlsx(header: list[str], body: list[list[object]]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Отчёт")
    sheet.append(header)
    for row in body:
        sheet.append(list(row))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _xls(header: list[str], body: list[list[object]]) -> bytes:
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Отчёт")
    for row_index, row in enumerate([header, *body]):
        for column_index, value in enumerate(row):
            sheet.write(row_index, column_index, "" if value is None else value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _font_path(filename: str) -> Path:
    candidates = [
        BACKEND_DIR / "app" / "assets" / filename,
        PROJECT_DIR / "frontend" / "design" / "font" / filename,
    ]
    path = next((candidate for candidate in candidates if candidate.is_file()), None)
    if path is None:
        raise APIError(500, "report_font_missing", "Не найден фирменный шрифт для PDF-отчёта.")
    return path


def _pdf(header: list[str], body: list[list[object]], title: str, summary: str) -> bytes:
    pdf = FPDF(orientation="L", format="A4")
    pdf.set_margins(10, 10, 10)
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.add_font("RostelecomBasis", fname=str(_font_path("RostelecomBasis-Regular.otf")))
    pdf.add_font("RostelecomBasis", style="B", fname=str(_font_path("RostelecomBasis-Bold.otf")))
    pdf.add_page()
    pdf.set_text_color(29, 29, 34)
    pdf.set_font("RostelecomBasis", style="B", size=16)
    pdf.multi_cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(86, 86, 95)
    pdf.set_font("RostelecomBasis", size=9)
    pdf.multi_cell(0, 5, summary, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_text_color(29, 29, 34)
    pdf.set_font("RostelecomBasis", size=7)
    rows = [[str(value or "") for value in row] for row in [header, *body]]
    pdf.table(
        rows=rows,
        line_height=5,
        headings_style=FontFace(emphasis="BOLD", fill_color=(240, 224, 255)),
    )
    return bytes(pdf.output())


def create_report(format_name: str, header: list[str], body: list[list[object]], *, title: str, summary: str) -> tuple[bytes, str]:
    if not header:
        raise APIError(422, "report_columns_required", "Выберите хотя бы одну колонку.")
    if not body:
        raise APIError(422, "report_rows_required", "В отчёте нет строк.")
    if format_name == "xlsx":
        data = _xlsx(header, body)
    elif format_name == "xls":
        data = _xls(header, body)
    elif format_name == "pdf":
        data = _pdf(header, body, title, summary)
    else:
        raise APIError(415, "unsupported_report_format", "Неподдерживаемый формат отчёта.")
    return data, CONTENT_TYPES[format_name]
